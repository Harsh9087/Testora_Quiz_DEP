import io
import re
import secrets
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, Response, status
from sqlalchemy.orm import Session
from sqlalchemy import or_

from app.database import get_db
from app.models.user import User
from app.models.workspace import Workspace
from app.models.registration_form import StudentRegistrationForm, StudentFormSubmission
from app.models.student_directory import StudentDirectory, DirectoryStudent
from app.schemas.registration_form import (
    StudentRegistrationFormCreate,
    StudentRegistrationFormUpdate,
    StudentRegistrationFormResponse,
    StudentRegistrationFormDetailResponse,
    PublicFormMetaResponse,
    StudentFormSubmissionCreate,
    StudentFormSubmissionResponse,
    ImportToDirectoryRequest,
    ImportToDirectoryResponse
)
from app.utils.security import get_current_user
from app.services.workspace_service import get_current_workspace

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/registration-forms", tags=["registration_forms"])

def _generate_unique_slug(db: Session) -> str:
    """Generates a clean, human-friendly 8-character unique alphanumeric slug."""
    for _ in range(10):
        # Generate 8 random alphanumeric chars
        token = secrets.token_hex(4)
        exists = db.query(StudentRegistrationForm).filter(
            StudentRegistrationForm.share_slug == token
        ).first()
        if not exists:
            return token
    # Fallback to longer urlsafe token
    return secrets.token_urlsafe(8).replace("-", "").replace("_", "")[:10]

# ══════════════════════════════════════════════════════════════════════════
# PUBLIC ENDPOINTS (No login required for students)
# ══════════════════════════════════════════════════════════════════════════

@router.get("/public/{share_slug}", response_model=PublicFormMetaResponse)
def get_public_form_meta(share_slug: str, db: Session = Depends(get_db)):
    """Public endpoint to fetch registration form metadata for students."""
    form = db.query(StudentRegistrationForm).filter(
        StudentRegistrationForm.share_slug == share_slug,
        StudentRegistrationForm.is_deleted == False
    ).first()

    if not form:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Registration form not found or has expired."
        )

    return {
        "title": form.title,
        "description": form.description,
        "share_slug": form.share_slug,
        "is_active": form.is_active,
        "created_at": form.created_at
    }

@router.post("/public/{share_slug}/submit", response_model=StudentFormSubmissionResponse, status_code=status.HTTP_201_CREATED)
def submit_student_registration(
    share_slug: str,
    payload: StudentFormSubmissionCreate,
    db: Session = Depends(get_db)
):
    """Public endpoint for students to submit the 6 mandatory registration fields."""
    form = db.query(StudentRegistrationForm).filter(
        StudentRegistrationForm.share_slug == share_slug,
        StudentRegistrationForm.is_deleted == False
    ).first()

    if not form:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Registration form not found."
        )

    if not form.is_active:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="This registration form is currently closed for new submissions by the instructor."
        )

    clean_email = payload.email.strip().lower()
    clean_roll = payload.roll_number.strip().upper()
    clean_name = payload.full_name.strip()
    clean_phone = payload.contact_number.strip()
    clean_section = payload.section.strip()
    clean_branch = payload.branch.strip()

    # Check for duplicate submission within this form
    existing_sub = db.query(StudentFormSubmission).filter(
        StudentFormSubmission.form_id == form.id,
        StudentFormSubmission.is_deleted == False,
        or_(
            StudentFormSubmission.email == clean_email,
            StudentFormSubmission.roll_number == clean_roll
        )
    ).first()

    if existing_sub:
        dup_field = "Roll Number" if existing_sub.roll_number == clean_roll else "Email Address"
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"A student response with this {dup_field} ({clean_roll if dup_field == 'Roll Number' else clean_email}) has already been registered."
        )

    submission = StudentFormSubmission(
        form_id=form.id,
        full_name=clean_name,
        email=clean_email,
        roll_number=clean_roll,
        contact_number=clean_phone,
        section=clean_section,
        branch=clean_branch,
        status="submitted"
    )

    db.add(submission)
    db.commit()
    db.refresh(submission)
    return submission

# ══════════════════════════════════════════════════════════════════════════
# FACULTY AUTHENTICATED ENDPOINTS
# ══════════════════════════════════════════════════════════════════════════

def _get_authorized_form(form_id: str, workspace: Workspace, db: Session) -> StudentRegistrationForm:
    form = db.query(StudentRegistrationForm).filter(
        StudentRegistrationForm.id == form_id,
        StudentRegistrationForm.workspace_id == workspace.id,
        StudentRegistrationForm.is_deleted == False
    ).first()
    if not form:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Registration form not found or does not belong to active workspace."
        )
    return form

@router.get("", response_model=List[StudentRegistrationFormResponse])
@router.get("/", response_model=List[StudentRegistrationFormResponse])
def list_faculty_registration_forms(
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """List all registration forms for the current workspace with response counts."""
    forms = db.query(StudentRegistrationForm).filter(
        StudentRegistrationForm.workspace_id == current_workspace.id,
        StudentRegistrationForm.is_deleted == False
    ).order_by(StudentRegistrationForm.created_at.desc()).all()

    results = []
    for f in forms:
        count = db.query(StudentFormSubmission).filter(
            StudentFormSubmission.form_id == f.id,
            StudentFormSubmission.is_deleted == False
        ).count()
        results.append({
            "id": f.id,
            "workspace_id": f.workspace_id,
            "created_by": f.created_by,
            "title": f.title,
            "description": f.description,
            "share_slug": f.share_slug,
            "is_active": f.is_active,
            "submission_count": count,
            "created_at": f.created_at
        })
    return results

@router.post("", response_model=StudentRegistrationFormResponse, status_code=status.HTTP_201_CREATED)
@router.post("/", response_model=StudentRegistrationFormResponse, status_code=status.HTTP_201_CREATED)
def create_registration_form(
    payload: StudentRegistrationFormCreate,
    current_user: User = Depends(get_current_user),
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """Faculty creates a new student data collection form by entering just the form name."""
    slug = _generate_unique_slug(db)
    form = StudentRegistrationForm(
        workspace_id=current_workspace.id,
        created_by=current_user.id,
        title=payload.title.strip(),
        description=payload.description.strip() if payload.description else None,
        share_slug=slug,
        is_active=True
    )
    db.add(form)
    db.commit()
    db.refresh(form)
    return {
        "id": form.id,
        "workspace_id": form.workspace_id,
        "created_by": form.created_by,
        "title": form.title,
        "description": form.description,
        "share_slug": form.share_slug,
        "is_active": form.is_active,
        "submission_count": 0,
        "created_at": form.created_at
    }

@router.get("/{form_id}", response_model=StudentRegistrationFormDetailResponse)
def get_registration_form_detail(
    form_id: str,
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """Fetch details and all student submissions for a specific form."""
    form = _get_authorized_form(form_id, current_workspace, db)
    submissions = db.query(StudentFormSubmission).filter(
        StudentFormSubmission.form_id == form.id,
        StudentFormSubmission.is_deleted == False
    ).order_by(StudentFormSubmission.created_at.desc()).all()

    return {
        "id": form.id,
        "workspace_id": form.workspace_id,
        "created_by": form.created_by,
        "title": form.title,
        "description": form.description,
        "share_slug": form.share_slug,
        "is_active": form.is_active,
        "submission_count": len(submissions),
        "created_at": form.created_at,
        "submissions": submissions
    }

@router.get("/{form_id}/submissions", response_model=List[StudentFormSubmissionResponse])
def get_registration_form_submissions(
    form_id: str,
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """Fetch all student submissions for a specific registration form."""
    form = _get_authorized_form(form_id, current_workspace, db)
    submissions = db.query(StudentFormSubmission).filter(
        StudentFormSubmission.form_id == form.id,
        StudentFormSubmission.is_deleted == False
    ).order_by(StudentFormSubmission.created_at.desc()).all()
    return submissions

@router.patch("/{form_id}", response_model=StudentRegistrationFormResponse)
def update_registration_form(
    form_id: str,
    payload: StudentRegistrationFormUpdate,
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """Update title, description, or active status (open/close) of a form."""
    form = _get_authorized_form(form_id, current_workspace, db)
    if payload.title is not None:
        setattr(form, "title", payload.title.strip())
    if payload.description is not None:
        setattr(form, "description", payload.description.strip())
    if payload.is_active is not None:
        setattr(form, "is_active", payload.is_active)

    db.commit()
    db.refresh(form)

    count = db.query(StudentFormSubmission).filter(
        StudentFormSubmission.form_id == form.id,
        StudentFormSubmission.is_deleted == False
    ).count()

    return {
        "id": form.id,
        "workspace_id": form.workspace_id,
        "created_by": form.created_by,
        "title": form.title,
        "description": form.description,
        "share_slug": form.share_slug,
        "is_active": form.is_active,
        "submission_count": count,
        "created_at": form.created_at
    }

@router.delete("/{form_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_registration_form(
    form_id: str,
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """Soft delete a registration form and its responses."""
    form = _get_authorized_form(form_id, current_workspace, db)
    form.delete()
    db.query(StudentFormSubmission).filter(
        StudentFormSubmission.form_id == form.id
    ).update({"is_deleted": True}, synchronize_session=False)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)

@router.delete("/{form_id}/submissions/{submission_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_form_submission(
    form_id: str,
    submission_id: str,
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """Delete an individual student response from the form."""
    form = _get_authorized_form(form_id, current_workspace, db)
    sub = db.query(StudentFormSubmission).filter(
        StudentFormSubmission.id == submission_id,
        StudentFormSubmission.form_id == form.id,
        StudentFormSubmission.is_deleted == False
    ).first()

    if not sub:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Submission record not found."
        )

    sub.delete()
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)

# ══════════════════════════════════════════════════════════════════════════
# DIRECTORY INTEGRATION & EXCEL EXPORT
# ══════════════════════════════════════════════════════════════════════════

@router.post("/{form_id}/import-to-directory", response_model=ImportToDirectoryResponse)
def import_submissions_to_directory(
    form_id: str,
    payload: ImportToDirectoryRequest,
    current_user: User = Depends(get_current_user),
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """
    Imports collected student submissions into an existing or newly created Student Directory.
    Validates duplicate roll numbers or emails against directory students.
    """
    form = _get_authorized_form(form_id, current_workspace, db)

    # 1. Resolve or create target StudentDirectory
    target_directory: Optional[StudentDirectory] = None
    if payload.target_mode == "existing":
        if not payload.directory_id:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Target directory ID is required when target_mode is 'existing'."
            )
        target_directory = db.query(StudentDirectory).filter(
            StudentDirectory.id == payload.directory_id,
            StudentDirectory.workspace_id == current_workspace.id,
            StudentDirectory.is_deleted == False
        ).first()
        if not target_directory:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Selected student directory not found in current workspace."
            )
    else:
        # Create new directory from form name or specified name
        dir_name = (payload.new_directory_name or form.title or "Imported Directory").strip()
        dir_desc = (payload.new_directory_description or f"Created via Data Collection Form: {form.title}").strip()
        target_directory = StudentDirectory(
            workspace_id=current_workspace.id,
            name=dir_name,
            description=dir_desc,
            created_by=current_user.id,
            is_active=True
        )
        db.add(target_directory)
        db.flush()

    # 2. Fetch candidate submissions
    query = db.query(StudentFormSubmission).filter(
        StudentFormSubmission.form_id == form.id,
        StudentFormSubmission.is_deleted == False
    )
    if payload.submission_ids:
        query = query.filter(StudentFormSubmission.id.in_(payload.submission_ids))

    submissions = query.all()
    if not submissions:
        return {
            "directory_id": target_directory.id,
            "directory_name": target_directory.name,
            "imported_count": 0,
            "skipped_duplicate_count": 0,
            "total_selected": 0,
            "messages": ["No submissions were selected for import."]
        }

    # 3. Fetch existing students in the directory to check for duplicates
    existing_directory_students = db.query(DirectoryStudent).filter(
        DirectoryStudent.directory_id == target_directory.id,
        DirectoryStudent.is_deleted == False
    ).all()

    existing_emails = {s.email.lower().strip() for s in existing_directory_students if s.email}
    existing_rolls = {s.roll_number.upper().strip() for s in existing_directory_students if s.roll_number}

    imported_count = 0
    skipped_count = 0
    messages = []

    for sub in submissions:
        sub_email = sub.email.lower().strip()
        sub_roll = sub.roll_number.upper().strip()

        if sub_email in existing_emails or sub_roll in existing_rolls:
            skipped_count += 1
            messages.append(f"Skipped duplicate: {sub.full_name} ({sub.roll_number} / {sub.email})")
            continue

        new_student = DirectoryStudent(
            directory_id=target_directory.id,
            name=sub.full_name.strip(),
            email=sub_email,
            roll_number=sub_roll,
            phone=sub.contact_number.strip(),
            division=sub.section.strip(),
            department=sub.branch.strip(),
            status="active"
        )
        db.add(new_student)
        existing_emails.add(sub_email)
        existing_rolls.add(sub_roll)

        sub.status = "imported"
        sub.imported_directory_id = target_directory.id
        imported_count += 1

    db.commit()

    return {
        "directory_id": target_directory.id,
        "directory_name": target_directory.name,
        "imported_count": imported_count,
        "skipped_duplicate_count": skipped_count,
        "total_selected": len(submissions),
        "messages": messages
    }

@router.get("/{form_id}/export-excel")
def export_submissions_excel(
    form_id: str,
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """
    Exports collected student responses to an Excel (.xlsx) file
    formatted compatibly with TestoraQuiz's Student Roster template.
    """
    form = _get_authorized_form(form_id, current_workspace, db)
    submissions = db.query(StudentFormSubmission).filter(
        StudentFormSubmission.form_id == form.id,
        StudentFormSubmission.is_deleted == False
    ).order_by(StudentFormSubmission.created_at.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet("Student Responses")
    else:
        ws.title = "Student Responses"

    # Brand color styles
    header_fill = PatternFill(start_color="C84B18", end_color="C84B18", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10, color="242321")
    thin_border = Border(
        left=Side(style="thin", color="E5E0D8"),
        right=Side(style="thin", color="E5E0D8"),
        top=Side(style="thin", color="E5E0D8"),
        bottom=Side(style="thin", color="E5E0D8")
    )

    headers = [
        "Full Name",
        "Email Address",
        "Roll Number",
        "Phone Number",
        "Division / Section",
        "Department / Branch",
        "Submission Status",
        "Submitted At"
    ]
    ws.append(headers)

    # Style Header Row
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Add Data Rows
    for row_idx, sub in enumerate(submissions, start=2):
        row_values = [
            sub.full_name,
            sub.email,
            sub.roll_number,
            sub.contact_number,
            sub.section,
            sub.branch,
            sub.status.capitalize(),
            sub.created_at.strftime("%Y-%m-%d %H:%M") if sub.created_at else ""
        ]
        ws.append(row_values)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [3, 4, 5, 7, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-fit Column Widths with padding
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 18)

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    excel_bytes = buf.getvalue()

    safe_title = re.sub(r'[^a-zA-Z0-9_\-]', '_', form.title)
    filename = f"{safe_title}_responses.xlsx"

    return Response(
        content=excel_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
