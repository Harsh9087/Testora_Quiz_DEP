import csv
import io
import json
import re
import secrets
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, Response, status
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.user import User
from app.models.workspace import Workspace
from app.models.custom_form import CustomForm, CustomFormSubmission
from app.schemas.custom_form import (
    GenerateCustomFormRequest,
    GeneratedFormPreviewResponse,
    CustomFormCreate,
    CustomFormUpdate,
    CustomFormResponse,
    CustomFormDetailResponse,
    PublicCustomFormMetaResponse,
    CustomFormSubmissionCreate,
    CustomFormSubmissionResponse
)
from app.services.ai_form_service import ai_form_service
from app.utils.security import get_current_user
from app.services.workspace_service import get_current_workspace

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/custom-forms", tags=["custom_forms"])

def _generate_unique_custom_slug(db: Session) -> str:
    for _ in range(10):
        token = secrets.token_hex(4)
        exists = db.query(CustomForm).filter(CustomForm.share_slug == token).first()
        if not exists:
            return token
    return secrets.token_urlsafe(8).replace("-", "").replace("_", "")[:10]

def _get_authorized_custom_form(form_id: str, workspace: Workspace, db: Session) -> CustomForm:
    form = db.query(CustomForm).filter(
        CustomForm.id == form_id,
        CustomForm.workspace_id == workspace.id,
        CustomForm.is_deleted == False
    ).first()
    if not form:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Custom form not found or does not belong to active workspace."
        )
    return form

# ══════════════════════════════════════════════════════════════════════════
# PUBLIC ENDPOINTS FOR CUSTOM FORMS (No auth required)
# ══════════════════════════════════════════════════════════════════════════

@router.get("/public/{share_slug}", response_model=PublicCustomFormMetaResponse)
def get_public_custom_form(share_slug: str, db: Session = Depends(get_db)):
    """Public endpoint to fetch published custom form structure for respondents."""
    form = db.query(CustomForm).filter(
        CustomForm.share_slug == share_slug,
        CustomForm.is_deleted == False
    ).first()

    if not form:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Custom form not found."
        )

    if not form.is_published:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="This form is currently a draft and has not been published yet."
        )

    try:
        fields = json.loads(form.fields_schema) if form.fields_schema else []
    except Exception:
        fields = []

    return {
        "title": form.title,
        "description": form.description,
        "fields": fields,
        "share_slug": form.share_slug,
        "is_active": form.is_active,
        "is_published": form.is_published,
        "created_at": form.created_at
    }

@router.post("/public/{share_slug}/submit", response_model=CustomFormSubmissionResponse, status_code=status.HTTP_201_CREATED)
def submit_custom_form(
    share_slug: str,
    payload: CustomFormSubmissionCreate,
    db: Session = Depends(get_db)
):
    """Public endpoint to submit answers to a custom form."""
    form = db.query(CustomForm).filter(
        CustomForm.share_slug == share_slug,
        CustomForm.is_deleted == False
    ).first()

    if not form:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Custom form not found."
        )

    if not form.is_published:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="This form is not open for public submissions."
        )

    if not form.is_active:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="This form is currently closed for responses."
        )

    # Validate required fields
    try:
        fields = json.loads(form.fields_schema) if form.fields_schema else []
    except Exception:
        fields = []

    for f in fields:
        fid = f.get("id")
        if f.get("required"):
            val = payload.responses.get(fid)
            if val is None or (isinstance(val, str) and not val.strip()) or (isinstance(val, list) and len(val) == 0):
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail=f"Field '{f.get('label')}' is required."
                )

    # Identify respondent (e.g. from email/name field if found)
    identifier = None
    for f in fields:
        fid = f.get("id")
        ftype = f.get("field_type")
        flabel = f.get("label", "").lower()
        if ftype == "email" or "email" in flabel:
            val = payload.responses.get(fid)
            if val:
                identifier = str(val).strip()
                break
        elif "name" in flabel and not identifier:
            val = payload.responses.get(fid)
            if val:
                identifier = str(val).strip()

    submission = CustomFormSubmission(
        form_id=form.id,
        responses_json=json.dumps(payload.responses),
        respondent_identifier=identifier,
        status="submitted"
    )
    db.add(submission)
    db.commit()
    db.refresh(submission)

    return {
        "id": submission.id,
        "form_id": submission.form_id,
        "responses": payload.responses,
        "respondent_identifier": submission.respondent_identifier,
        "created_at": submission.created_at,
        "status": submission.status
    }

# ══════════════════════════════════════════════════════════════════════════
# FACULTY AUTHENTICATED ENDPOINTS
# ══════════════════════════════════════════════════════════════════════════

@router.post("/generate", response_model=GeneratedFormPreviewResponse)
def generate_custom_form_preview(
    payload: GenerateCustomFormRequest,
    current_user: User = Depends(get_current_user),
    current_workspace: Workspace = Depends(get_current_workspace)
):
    """Synthesizes structured form fields from a natural language prompt using Gemini."""
    result = ai_form_service.generate_form(payload.prompt)
    return result

@router.post("", response_model=CustomFormResponse, status_code=status.HTTP_201_CREATED)
@router.post("/", response_model=CustomFormResponse, status_code=status.HTTP_201_CREATED)
def create_custom_form(
    payload: CustomFormCreate,
    current_user: User = Depends(get_current_user),
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """Saves a new custom form with fields."""
    slug = _generate_unique_custom_slug(db)
    fields_data = [f.model_dump() for f in payload.fields]

    form = CustomForm(
        workspace_id=current_workspace.id,
        created_by=current_user.id,
        title=payload.title.strip(),
        description=payload.description.strip() if payload.description else None,
        prompt=payload.prompt.strip() if payload.prompt else None,
        fields_schema=json.dumps(fields_data),
        share_slug=slug,
        is_published=payload.is_published,
        is_active=True
    )
    db.add(form)
    db.commit()
    db.refresh(form)

    return {
        "id": form.id,
        "workspace_id": form.workspace_id,
        "created_by": form.created_by,
        "title": form.title,
        "description": form.description,
        "prompt": form.prompt,
        "fields": fields_data,
        "share_slug": form.share_slug,
        "is_published": form.is_published,
        "is_active": form.is_active,
        "submission_count": 0,
        "created_at": form.created_at
    }

@router.get("", response_model=List[CustomFormResponse])
@router.get("/", response_model=List[CustomFormResponse])
def list_custom_forms(
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """List all custom forms in the current workspace with submission counts."""
    forms = db.query(CustomForm).filter(
        CustomForm.workspace_id == current_workspace.id,
        CustomForm.is_deleted == False
    ).order_by(CustomForm.created_at.desc()).all()

    results = []
    for f in forms:
        count = db.query(CustomFormSubmission).filter(
            CustomFormSubmission.form_id == f.id,
            CustomFormSubmission.is_deleted == False
        ).count()
        try:
            fields = json.loads(f.fields_schema) if f.fields_schema else []
        except Exception:
            fields = []

        results.append({
            "id": f.id,
            "workspace_id": f.workspace_id,
            "created_by": f.created_by,
            "title": f.title,
            "description": f.description,
            "prompt": f.prompt,
            "fields": fields,
            "share_slug": f.share_slug,
            "is_published": f.is_published,
            "is_active": f.is_active,
            "submission_count": count,
            "created_at": f.created_at
        })
    return results

@router.get("/{form_id}", response_model=CustomFormDetailResponse)
def get_custom_form_detail(
    form_id: str,
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """Get full details of a custom form and all its submissions."""
    form = _get_authorized_custom_form(form_id, current_workspace, db)
    submissions = db.query(CustomFormSubmission).filter(
        CustomFormSubmission.form_id == form.id,
        CustomFormSubmission.is_deleted == False
    ).order_by(CustomFormSubmission.created_at.desc()).all()

    try:
        fields = json.loads(form.fields_schema) if form.fields_schema else []
    except Exception:
        fields = []

    parsed_subs = []
    for s in submissions:
        try:
            resp_dict = json.loads(s.responses_json) if s.responses_json else {}
        except Exception:
            resp_dict = {}
        parsed_subs.append({
            "id": s.id,
            "form_id": s.form_id,
            "responses": resp_dict,
            "respondent_identifier": s.respondent_identifier,
            "created_at": s.created_at,
            "status": s.status
        })

    return {
        "id": form.id,
        "workspace_id": form.workspace_id,
        "created_by": form.created_by,
        "title": form.title,
        "description": form.description,
        "prompt": form.prompt,
        "fields": fields,
        "share_slug": form.share_slug,
        "is_published": form.is_published,
        "is_active": form.is_active,
        "submission_count": len(submissions),
        "created_at": form.created_at,
        "submissions": parsed_subs
    }

@router.get("/{form_id}/submissions", response_model=List[CustomFormSubmissionResponse])
def get_custom_form_submissions(
    form_id: str,
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """Fetch all submissions for a custom form."""
    form = _get_authorized_custom_form(form_id, current_workspace, db)
    submissions = db.query(CustomFormSubmission).filter(
        CustomFormSubmission.form_id == form.id,
        CustomFormSubmission.is_deleted == False
    ).order_by(CustomFormSubmission.created_at.desc()).all()

    parsed_subs = []
    for s in submissions:
        try:
            resp_dict = json.loads(s.responses_json) if s.responses_json else {}
        except Exception:
            resp_dict = {}
        parsed_subs.append({
            "id": s.id,
            "form_id": s.form_id,
            "responses": resp_dict,
            "respondent_identifier": s.respondent_identifier,
            "status": s.status or "submitted",
            "created_at": s.created_at,
            "submitted_at": s.created_at
        })
    return parsed_subs

@router.patch("/{form_id}", response_model=CustomFormResponse)
def update_custom_form(
    form_id: str,
    payload: CustomFormUpdate,
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """Update custom form properties or fields."""
    form = _get_authorized_custom_form(form_id, current_workspace, db)

    if payload.title is not None:
        setattr(form, "title", payload.title.strip())
    if payload.description is not None:
        setattr(form, "description", payload.description.strip())
    if payload.fields is not None:
        setattr(form, "fields_schema", json.dumps([f.model_dump() for f in payload.fields]))
    if payload.is_published is not None:
        setattr(form, "is_published", payload.is_published)
    if payload.is_active is not None:
        setattr(form, "is_active", payload.is_active)

    db.commit()
    db.refresh(form)

    count = db.query(CustomFormSubmission).filter(
        CustomFormSubmission.form_id == form.id,
        CustomFormSubmission.is_deleted == False
    ).count()

    try:
        fields = json.loads(form.fields_schema) if form.fields_schema else []
    except Exception:
        fields = []

    return {
        "id": form.id,
        "workspace_id": form.workspace_id,
        "created_by": form.created_by,
        "title": form.title,
        "description": form.description,
        "prompt": form.prompt,
        "fields": fields,
        "share_slug": form.share_slug,
        "is_published": form.is_published,
        "is_active": form.is_active,
        "submission_count": count,
        "created_at": form.created_at
    }

@router.delete("/{form_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_custom_form(
    form_id: str,
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """Soft delete custom form and its responses."""
    form = _get_authorized_custom_form(form_id, current_workspace, db)
    form.delete()
    db.query(CustomFormSubmission).filter(
        CustomFormSubmission.form_id == form.id
    ).update({"is_deleted": True}, synchronize_session=False)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)

@router.delete("/{form_id}/submissions/{submission_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_custom_form_submission(
    form_id: str,
    submission_id: str,
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """Delete an individual submission from the custom form."""
    form = _get_authorized_custom_form(form_id, current_workspace, db)
    sub = db.query(CustomFormSubmission).filter(
        CustomFormSubmission.id == submission_id,
        CustomFormSubmission.form_id == form.id,
        CustomFormSubmission.is_deleted == False
    ).first()

    if not sub:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Submission not found."
        )

    sub.delete()
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)

@router.get("/{form_id}/export-excel")
def export_custom_form_excel(
    form_id: str,
    current_workspace: Workspace = Depends(get_current_workspace),
    db: Session = Depends(get_db)
):
    """Exports custom form submissions to an Excel file with dynamic columns matching form fields."""
    form = _get_authorized_custom_form(form_id, current_workspace, db)
    submissions = db.query(CustomFormSubmission).filter(
        CustomFormSubmission.form_id == form.id,
        CustomFormSubmission.is_deleted == False
    ).order_by(CustomFormSubmission.created_at.asc()).all()

    try:
        fields = json.loads(form.fields_schema) if form.fields_schema else []
    except Exception:
        fields = []

    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet("Form Responses")
    else:
        ws.title = "Form Responses"

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10, color="242321")
    thin_border = Border(
        left=Side(style="thin", color="E5E0D8"),
        right=Side(style="thin", color="E5E0D8"),
        top=Side(style="thin", color="E5E0D8"),
        bottom=Side(style="thin", color="E5E0D8")
    )

    headers = ["#", "Submission Date", "Respondent"]
    field_id_map = []
    for f in fields:
        headers.append(f.get("label", "Question"))
        field_id_map.append(f.get("id"))

    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, sub in enumerate(submissions, start=2):
        try:
            resp_dict = json.loads(sub.responses_json) if sub.responses_json else {}
        except Exception:
            resp_dict = {}

        date_str = sub.created_at.strftime("%Y-%m-%d %H:%M") if sub.created_at else ""
        row_vals = [row_idx - 1, date_str, sub.respondent_identifier or "Anonymous"]

        for fid in field_id_map:
            val = resp_dict.get(fid, "")
            if isinstance(val, list):
                val = ", ".join(str(v) for v in val)
            row_vals.append(str(val) if val is not None else "")

        ws.append(row_vals)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [1, 2]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    excel_bytes = buf.getvalue()

    safe_title = re.sub(r'[^a-zA-Z0-9_\-]', '_', form.title)
    filename = f"{safe_title}_responses.xlsx"

    return Response(
        content=excel_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
